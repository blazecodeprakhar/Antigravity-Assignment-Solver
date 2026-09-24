"""
Antigravity Universal Academic Assignment Solver Engine v2.0
Multi-Language, Multi-Theme, Zero-Box Formatting Pipeline
"""

import os
import glob
import json
import re
import subprocess
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

# Theme Palette Specifications
THEMES = {
    "cmd_light": {
        "bg": "#FFFFFF",
        "title_bg": "#F0F3F6",
        "title_text": "#2C3E50",
        "border": "#CFD6DD",
        "prompt_text": "#004085",
        "output_text": "#111111"
    },
    "powershell_light": {
        "bg": "#F4F7FB",
        "title_bg": "#E2E8F0",
        "title_text": "#1E293B",
        "border": "#CBD5E1",
        "prompt_text": "#0284C7",
        "output_text": "#0F172A"
    },
    "macos_light": {
        "bg": "#FFFFFF",
        "title_bg": "#E9EAEB",
        "title_text": "#4A4A4A",
        "border": "#D1D5DB",
        "prompt_text": "#0D6EFD",
        "output_text": "#212529"
    }
}

def load_config():
    config_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "config.json")
    if os.path.exists(config_path):
        try:
            with open(config_path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return {
        "theme": "cmd_light",
        "font_size": 10.5,
        "line_spacing": 1.15,
        "strip_comments": True,
        "screenshot": {"width_inches": 5.8}
    }

def render_styled_terminal(command_line, output_text, output_png_path, title_str="Command Prompt - Execution", theme_name="cmd_light"):
    theme = THEMES.get(theme_name, THEMES["cmd_light"])
    
    font_path = "C:/Windows/Fonts/consola.ttf"
    font_bold_path = "C:/Windows/Fonts/consolab.ttf"
    header_font_path = "C:/Windows/Fonts/segoeui.ttf"
    
    font_size = 15
    font = ImageFont.truetype(font_path, font_size)
    font_bold = ImageFont.truetype(font_bold_path, font_size)
    header_font = ImageFont.truetype(header_font_path, 13)
    
    lines = [command_line] + output_text.split('\n')
    
    padding_x = 25
    padding_y = 18
    title_bar_height = 38
    line_height = 24
    
    max_line_width = max(font.getlength(line) for line in lines)
    img_width = int(max_line_width + padding_x * 2 + 30)
    img_width = max(img_width, 680)
    
    img_height = title_bar_height + padding_y * 2 + len(lines) * line_height + 12
    
    img = Image.new('RGB', (img_width, img_height), color=theme["bg"])
    draw = ImageDraw.Draw(img)
    
    # Titlebar
    draw.rectangle([0, 0, img_width - 1, title_bar_height], fill=theme["title_bg"], outline=theme["border"])
    draw.text((16, 11), title_str, fill=theme["title_text"], font=header_font)
    
    # Window Buttons (Close, Maximize, Minimize)
    btn_w = 46
    draw.rectangle([img_width - btn_w, 0, img_width - 1, title_bar_height], fill='#E81123')
    draw.line([(img_width - btn_w + 18, 14), (img_width - btn_w + 28, 24)], fill='#FFFFFF', width=1)
    draw.line([(img_width - btn_w + 18, 24), (img_width - btn_w + 28, 14)], fill='#FFFFFF', width=1)
    
    draw.rectangle([img_width - btn_w*2, 0, img_width - btn_w - 1, title_bar_height], fill=theme["title_bg"])
    draw.rectangle([img_width - btn_w*2 + 18, 14, img_width - btn_w*2 + 28, 24], outline='#333333', width=1)
    
    draw.rectangle([img_width - btn_w*3, 0, img_width - btn_w*2 - 1, title_bar_height], fill=theme["title_bg"])
    draw.line([(img_width - btn_w*3 + 18, 19), (img_width - btn_w*3 + 28, 19)], fill='#333333', width=1)
    
    # Outer Border & Content Area
    draw.rectangle([0, title_bar_height, img_width - 1, img_height - 1], outline=theme["border"], width=1)
    draw.rectangle([1, title_bar_height + 1, img_width - 2, img_height - 2], fill=theme["bg"])
    
    # Draw Lines
    y = title_bar_height + padding_y
    draw.text((padding_x, y), command_line, fill=theme["prompt_text"], font=font_bold)
    y += line_height
    
    for line in output_text.split('\n'):
        draw.text((padding_x, y), line, fill=theme["output_text"], font=font)
        y += line_height
        
    img.save(output_png_path, "PNG")

def strip_code_comments(code_str, ext):
    lines = code_str.split('\n')
    clean_lines = []
    for line in lines:
        stripped = line.strip()
        if ext in ['.py']:
            if stripped.startswith('#'):
                continue
            if '#' in line:
                line = line.split('#')[0].rstrip()
        elif ext in ['.java', '.cpp', '.c', '.cs', '.js', '.go', '.rs']:
            if stripped.startswith('//') or stripped.startswith('/*') or stripped.startswith('*'):
                continue
            if '//' in line:
                line = line.split('//')[0].rstrip()
        clean_lines.append(line)
    return '\n'.join(clean_lines)

def run_code_file(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    base = os.path.basename(filepath)
    folder = os.path.dirname(filepath)
    
    if ext == '.py':
        cmd_prompt = f"C:\\Users\\prakh\\OneDrive\\Desktop\\ai class work\\workspace> python {base}"
        res = subprocess.run([sys.executable, filepath], capture_output=True, text=True, cwd=folder)
        return cmd_prompt, res.stdout.strip()
    elif ext == '.js':
        cmd_prompt = f"C:\\Users\\prakh\\OneDrive\\Desktop\\ai class work\\workspace> node {base}"
        res = subprocess.run(["node", filepath], capture_output=True, text=True, cwd=folder)
        return cmd_prompt, res.stdout.strip()
    elif ext == '.cpp':
        exe_name = os.path.splitext(base)[0] + ".exe"
        cmd_prompt = f"C:\\Users\\prakh\\OneDrive\\Desktop\\ai class work\\workspace> g++ {base} -o {exe_name} && {exe_name}"
        subprocess.run(["g++", base, "-o", exe_name], capture_output=True, text=True, cwd=folder)
        res = subprocess.run([os.path.join(folder, exe_name)], capture_output=True, text=True, cwd=folder)
        return cmd_prompt, res.stdout.strip()
    elif ext == '.c':
        exe_name = os.path.splitext(base)[0] + ".exe"
        cmd_prompt = f"C:\\Users\\prakh\\OneDrive\\Desktop\\ai class work\\workspace> gcc {base} -o {exe_name} && {exe_name}"
        subprocess.run(["gcc", base, "-o", exe_name], capture_output=True, text=True, cwd=folder)
        res = subprocess.run([os.path.join(folder, exe_name)], capture_output=True, text=True, cwd=folder)
        return cmd_prompt, res.stdout.strip()
    elif ext == '.java':
        cmd_prompt = f"C:\\Users\\prakh\\OneDrive\\Desktop\\ai class work\\workspace> javac {base} && java {os.path.splitext(base)[0]}"
        subprocess.run(["javac", base], capture_output=True, text=True, cwd=folder)
        res = subprocess.run(["java", os.path.splitext(base)[0]], capture_output=True, text=True, cwd=folder)
        return cmd_prompt, res.stdout.strip()
    else:
        cmd_prompt = f"C:\\Users\\prakh\\OneDrive\\Desktop\\ai class work\\workspace> python {base}"
        res = subprocess.run([sys.executable, filepath], capture_output=True, text=True, cwd=folder)
        return cmd_prompt, res.stdout.strip()

def process_assignment():
    cfg = load_config()
    workspace_root = os.path.dirname(os.path.abspath(__file__))
    raw_dir = os.path.join(workspace_root, "raw_file")
    workspace_dir = os.path.join(workspace_root, "workspace")
    output_dir = os.path.join(workspace_root, "final_output")
    
    os.makedirs(workspace_dir, exist_ok=True)
    os.makedirs(output_dir, exist_ok=True)
    
    raw_files = glob.glob(os.path.join(raw_dir, "*.docx"))
    if not raw_files:
        print("[!] Warning: No raw .docx file found in raw_file/. Place an assignment docx file there.")
        return
        
    raw_docx_path = raw_files[0]
    filename = os.path.basename(raw_docx_path)
    output_docx_path = os.path.join(output_dir, filename)
    
    print("=" * 60)
    print(f"[+] {cfg.get('project_name', 'ANTIGRAVITY ASSIGNMENT SOLVER')}")
    print(f"    Version: {cfg.get('version', '2.0.0')}")
    print(f"    Target File: {filename}")
    print("=" * 60)
    
    code_files = sorted(glob.glob(os.path.join(workspace_dir, "Problem*.*")))
    code_files = [f for f in code_files if not f.endswith('.png') and not f.endswith('.exe') and not f.endswith('.class')]
    
    if not code_files:
        print("[!] Warning: No solution code files (Problem1.py, Problem2.java, etc.) found in workspace/!")
        return
        
    prob_data = []
    for idx, code_file in enumerate(code_files, 1):
        basename = os.path.splitext(os.path.basename(code_file))[0]
        ext = os.path.splitext(code_file)[1].lower()
        
        print(f"[+] Executing & rendering screenshot: {os.path.basename(code_file)}")
        cmd_str, out_text = run_code_file(code_file)
        
        png_path = os.path.join(workspace_dir, f"{basename}_output.png")
        render_styled_terminal(cmd_str, out_text, png_path, title_str=f"Command Prompt - {os.path.basename(code_file)}", theme_name=cfg.get("theme", "cmd_light"))
        
        prob_data.append({
            "num": idx,
            "code_file": code_file,
            "ext": ext,
            "png_path": png_path
        })
        
    print("[+] Appending solution sections to document...")
    doc = Document(raw_docx_path)
    if cfg.get("page_break_after_original", True):
        doc.add_page_break()
    
    for item in prob_data:
        code_file = item["code_file"]
        ext = item["ext"]
        png_path = item["png_path"]
        num = item["num"]
        
        # Query Heading
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(14)
        p_q.paragraph_format.space_after = Pt(4)
        r_q = p_q.add_run(f"Query {num}:")
        r_q.bold = True
        r_q.font.size = Pt(12)
        r_q.font.name = 'Calibri'
        
        with open(code_file, "r", encoding="utf-8") as f:
            raw_code = f.read()
            
        clean_code = strip_code_comments(raw_code, ext) if cfg.get("strip_comments", True) else raw_code
        code_lines = clean_code.strip().split('\n')
        
        # Format Code as PLAIN PARAGRAPHS (NO BOXES / NO TABLES)
        for line in code_lines:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = cfg.get("line_spacing", 1.15)
            run = p.add_run(line if line else " ")
            run.font.name = cfg.get("default_font", "Consolas")
            run.font.size = Pt(cfg.get("font_size", 10.5))
            run.font.color.rgb = RGBColor(0, 0, 0)
            
        # Output Heading
        p_out = doc.add_paragraph()
        p_out.paragraph_format.space_before = Pt(10)
        p_out.paragraph_format.space_after = Pt(4)
        r_out = p_out.add_run("Output:")
        r_out.bold = True
        r_out.font.size = Pt(12)
        r_out.font.name = 'Calibri'
        
        # Output Screenshot Image
        p_img = doc.add_paragraph()
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(14)
        r_img = p_img.add_run()
        width_inches = cfg.get("screenshot", {}).get("width_inches", 5.8)
        r_img.add_picture(png_path, width=Inches(width_inches))
        
    try:
        doc.save(output_docx_path)
        print(f"[SUCCESS] Solved document generated at: {output_docx_path}")
    except PermissionError:
        alt_path = os.path.join(output_dir, f"solved_{filename}")
        doc.save(alt_path)
        print(f"[SUCCESS] (File open in MS Word) Saved to: {alt_path}")

if __name__ == "__main__":
    process_assignment()
